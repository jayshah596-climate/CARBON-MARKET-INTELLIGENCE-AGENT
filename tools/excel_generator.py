"""
Excel Dashboard Generator for Carbon Market Intelligence Reports.

Produces a colour-coded .xlsx workbook with multiple sheets:
  - Carbon Dashboard  (KPI overview)
  - Compliance Markets
  - Voluntary Carbon Market
  - Carbon Offsets
  - Risk Analysis
  - Opportunities & Forecast
"""

from __future__ import annotations
import logging
from io import BytesIO
from pathlib import Path
from datetime import datetime
from typing import Any

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill,
    Font,
    Alignment,
    Border,
    Side,
    GradientFill,
)
from openpyxl.utils import get_column_letter

from models.carbon_models import CarbonIntelligenceReport

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Colour palette
# ---------------------------------------------------------------------------
COLORS = {
    "Green":       "00B050",
    "Yellow":      "FFC000",
    "Red":         "FF0000",
    "Header":      "1F3864",   # dark navy for header rows
    "SubHeader":   "2E75B6",   # medium blue for sub-headers
    "LightBlue":   "D9E1F2",   # light blue for alternating rows
    "White":       "FFFFFF",
    "DarkGrey":    "404040",
    "LightGrey":   "F2F2F2",
    "Gold":        "FFD700",
}


def _fill(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=hex_color)


def _font(bold: bool = False, color: str = "000000", size: int = 11) -> Font:
    return Font(bold=bold, color=color, size=size)


def _border() -> Border:
    thin = Side(style="thin", color="BFBFBF")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _align(horizontal: str = "left", wrap: bool = True) -> Alignment:
    return Alignment(horizontal=horizontal, vertical="center", wrap_text=wrap)


def _perf_color(performance: str) -> str:
    """Map performance string to hex color."""
    mapping = {
        "Good": COLORS["Green"],
        "Moderate": COLORS["Yellow"],
        "Risk": COLORS["Red"],
        "High": COLORS["Red"],
        "Medium": COLORS["Yellow"],
        "Low": COLORS["Green"],
        "Strong": COLORS["Green"],
        "Balanced": COLORS["Yellow"],
        "Weak": COLORS["Red"],
    }
    return mapping.get(performance, COLORS["Yellow"])


def _color_code_hex(color_code: str) -> str:
    return {
        "Green": COLORS["Green"],
        "Yellow": COLORS["Yellow"],
        "Red": COLORS["Red"],
    }.get(color_code, COLORS["Yellow"])


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------

class ExcelGenerator:
    """Generates a colour-coded Excel workbook from a CarbonIntelligenceReport."""

    def generate(self, report: CarbonIntelligenceReport, output_path: str | Path) -> str:
        """
        Build the workbook and save to *output_path*.

        Returns the absolute path to the saved file.
        """
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        # Remove default empty sheet
        wb.remove(wb.active)  # type: ignore[arg-type]

        self._add_dashboard_sheet(wb, report)
        self._add_compliance_sheet(wb, report)
        self._add_vcm_sheet(wb, report)
        self._add_offsets_sheet(wb, report)
        self._add_risk_sheet(wb, report)
        self._add_opportunities_sheet(wb, report)

        wb.save(str(output_path))
        logger.info("Excel workbook saved: %s", output_path)
        return str(output_path.resolve())

    def generate_bytes(self, report: CarbonIntelligenceReport) -> BytesIO:
        """Build the workbook and return it as a BytesIO buffer (for streaming)."""
        wb = Workbook()
        wb.remove(wb.active)  # type: ignore[arg-type]
        self._add_dashboard_sheet(wb, report)
        self._add_compliance_sheet(wb, report)
        self._add_vcm_sheet(wb, report)
        self._add_offsets_sheet(wb, report)
        self._add_risk_sheet(wb, report)
        self._add_opportunities_sheet(wb, report)
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    # ------------------------------------------------------------------
    # Sheet: Carbon Dashboard
    # ------------------------------------------------------------------

    def _add_dashboard_sheet(self, wb: Workbook, report: CarbonIntelligenceReport) -> None:
        ws = wb.create_sheet("Carbon Dashboard")
        ws.sheet_view.showGridLines = False

        # Title
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = "🌍 CARBON MARKET INTELLIGENCE DASHBOARD"
        title_cell.fill = _fill(COLORS["Header"])
        title_cell.font = Font(bold=True, color="FFFFFF", size=16)
        title_cell.alignment = _align("center", wrap=False)
        ws.row_dimensions[1].height = 35

        # Date subtitle
        ws.merge_cells("A2:G2")
        date_cell = ws["A2"]
        date_cell.value = f"Report Date: {report.word_report.date or datetime.now().strftime('%Y-%m-%d')}"
        date_cell.fill = _fill(COLORS["SubHeader"])
        date_cell.font = Font(bold=False, color="FFFFFF", size=11)
        date_cell.alignment = _align("center", wrap=False)
        ws.row_dimensions[2].height = 22

        # Market Overview section header
        row = 4
        ws.merge_cells(f"A{row}:G{row}")
        ws[f"A{row}"].value = "MARKET OVERVIEW"
        ws[f"A{row}"].fill = _fill(COLORS["SubHeader"])
        ws[f"A{row}"].font = Font(bold=True, color="FFFFFF", size=12)
        ws[f"A{row}"].alignment = _align("center", wrap=False)
        ws.row_dimensions[row].height = 22

        # Headers
        row += 1
        headers = ["Metric", "Value", "Currency / Unit", "Trend", "Performance", "Risk Level", "Signal"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=h)
            cell.fill = _fill(COLORS["DarkGrey"])
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.alignment = _align("center", wrap=False)
            cell.border = _border()
        ws.row_dimensions[row].height = 20

        # Market overview data
        ov = report.market_overview
        overview_rows = [
            ("Market Type", ov.market_type, "", "", "", "", ""),
            ("Region", ov.region, "", "", "", "", ""),
            (
                "Latest Carbon Price",
                ov.latest_price.value,
                f"{ov.latest_price.currency} {ov.latest_price.unit}",
                ov.price_trend,
                ov.performance_rating,
                ov.volatility,
                ov.color_code,
            ),
        ]
        for data_row in overview_rows:
            row += 1
            for col, val in enumerate(data_row, 1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.alignment = _align("left", wrap=True)
                cell.border = _border()
                if col == 7 and val in ("Green", "Yellow", "Red"):
                    cell.fill = _fill(_color_code_hex(val))
                    cell.value = f"● {val}"
                    cell.font = Font(bold=True, color="FFFFFF", size=10)
                elif col == 5 and val in ("Good", "Moderate", "Risk"):
                    cell.fill = _fill(_perf_color(val))
                    cell.font = Font(bold=True, color="FFFFFF", size=10)
            ws.row_dimensions[row].height = 18

        # KPI Summary from excel_analysis
        if report.excel_analysis:
            row += 2
            ws.merge_cells(f"A{row}:G{row}")
            ws[f"A{row}"].value = "KEY PERFORMANCE INDICATORS"
            ws[f"A{row}"].fill = _fill(COLORS["SubHeader"])
            ws[f"A{row}"].font = Font(bold=True, color="FFFFFF", size=12)
            ws[f"A{row}"].alignment = _align("center", wrap=False)

            row += 1
            kpi_headers = ["Sheet", "Metric", "Value", "Unit", "Performance", "", "Signal"]
            for col, h in enumerate(kpi_headers, 1):
                cell = ws.cell(row=row, column=col, value=h)
                cell.fill = _fill(COLORS["DarkGrey"])
                cell.font = Font(bold=True, color="FFFFFF", size=10)
                cell.alignment = _align("center", wrap=False)
                cell.border = _border()

            for i, kpi in enumerate(report.excel_analysis):
                row += 1
                bg = COLORS["LightBlue"] if i % 2 == 0 else COLORS["White"]
                vals = [kpi.sheet, kpi.metric, kpi.value, kpi.unit, kpi.performance, "", kpi.color_code]
                for col, val in enumerate(vals, 1):
                    cell = ws.cell(row=row, column=col, value=val)
                    cell.fill = _fill(bg)
                    cell.alignment = _align("left", wrap=True)
                    cell.border = _border()
                    if col == 5 and val in ("Good", "Moderate", "Risk"):
                        cell.fill = _fill(_perf_color(val))
                        cell.font = Font(bold=True, color="FFFFFF", size=10)
                    if col == 7 and val in ("Green", "Yellow", "Red"):
                        cell.fill = _fill(_color_code_hex(val))
                        cell.value = f"● {val}"
                        cell.font = Font(bold=True, color="FFFFFF", size=10)

        # Column widths
        col_widths = [22, 18, 20, 16, 14, 12, 12]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------
    # Sheet: Compliance Markets
    # ------------------------------------------------------------------

    def _add_compliance_sheet(self, wb: Workbook, report: CarbonIntelligenceReport) -> None:
        ws = wb.create_sheet("Compliance Markets")
        ws.sheet_view.showGridLines = False

        self._sheet_title(ws, "COMPLIANCE CARBON MARKETS — GLOBAL OVERVIEW", 7)

        headers = ["Market Name", "Region", "Price", "Currency", "Trend", "Performance", "Signal"]
        row = self._write_headers(ws, headers, start_row=3)

        if not report.compliance_markets:
            ws.cell(row=row + 1, column=1, value="No compliance market data available.")
        else:
            for i, mkt in enumerate(report.compliance_markets):
                row += 1
                bg = COLORS["LightBlue"] if i % 2 == 0 else COLORS["White"]
                vals = [
                    mkt.market_name, mkt.region, mkt.price, mkt.currency,
                    mkt.trend, mkt.performance, mkt.color_code,
                ]
                for col, val in enumerate(vals, 1):
                    cell = ws.cell(row=row, column=col, value=val)
                    cell.fill = _fill(bg)
                    cell.alignment = _align()
                    cell.border = _border()
                    if col == 6 and val in ("Good", "Moderate", "Risk"):
                        cell.fill = _fill(_perf_color(val))
                        cell.font = Font(bold=True, color="FFFFFF", size=10)
                    if col == 7 and val in ("Green", "Yellow", "Red"):
                        cell.fill = _fill(_color_code_hex(val))
                        cell.value = f"● {val}"
                        cell.font = Font(bold=True, color="FFFFFF", size=10)

        # Regulatory insights section
        if report.compliance_markets:
            row += 2
            ws.merge_cells(f"A{row}:G{row}")
            ws[f"A{row}"].value = "REGULATORY UPDATES"
            ws[f"A{row}"].fill = _fill(COLORS["SubHeader"])
            ws[f"A{row}"].font = Font(bold=True, color="FFFFFF", size=11)
            ws[f"A{row}"].alignment = _align("center")

            for mkt in report.compliance_markets:
                if mkt.regulatory_updates and mkt.regulatory_updates != "missing":
                    row += 1
                    ws.merge_cells(f"A{row}:G{row}")
                    ws[f"A{row}"].value = f"• [{mkt.market_name}] {mkt.regulatory_updates}"
                    ws[f"A{row}"].alignment = _align("left", wrap=True)
                    ws[f"A{row}"].fill = _fill(COLORS["LightGrey"])
                    ws[f"A{row}"].border = _border()
                    ws.row_dimensions[row].height = 30

        col_widths = [24, 16, 10, 10, 14, 14, 12]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------
    # Sheet: Voluntary Carbon Market
    # ------------------------------------------------------------------

    def _add_vcm_sheet(self, wb: Workbook, report: CarbonIntelligenceReport) -> None:
        ws = wb.create_sheet("Voluntary Carbon Market")
        ws.sheet_view.showGridLines = False

        self._sheet_title(ws, "VOLUNTARY CARBON MARKET (VCM) ANALYSIS", 4)

        vcm = report.voluntary_carbon_market
        row = 3

        kv_pairs = [
            ("Average Price", vcm.average_price, vcm.currency),
            ("Price Range", vcm.price_range, vcm.currency),
            ("Demand Trend", vcm.demand_trend, ""),
            ("Supply Trend", vcm.supply_trend, ""),
            ("Market Condition", vcm.market_condition, ""),
            ("Performance Signal", vcm.color_code, ""),
        ]
        headers = ["Indicator", "Value", "Unit / Currency", "Assessment"]
        self._write_headers(ws, headers, start_row=row, col_count=4)
        row += 1

        for i, (label, value, unit) in enumerate(kv_pairs):
            bg = COLORS["LightBlue"] if i % 2 == 0 else COLORS["White"]
            assessment = ""
            if label == "Market Condition":
                assessment = {"Strong": "✅ Strong demand", "Balanced": "⚖️ Balanced", "Weak": "⚠️ Oversupply risk"}.get(value, "")
            if label == "Performance Signal":
                assessment = value
                value = {"Green": "Positive", "Yellow": "Neutral", "Red": "Negative"}.get(value, value)

            for col, val in enumerate([label, value, unit, assessment], 1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.fill = _fill(bg)
                cell.alignment = _align()
                cell.border = _border()
                if col == 4 and assessment in ("Green", "Yellow", "Red"):
                    cell.fill = _fill(_color_code_hex(assessment))
                    cell.font = Font(bold=True, color="FFFFFF")
            row += 1

        # Key standards
        if vcm.key_standards:
            row += 1
            ws.merge_cells(f"A{row}:D{row}")
            ws[f"A{row}"].value = "KEY STANDARDS & REGISTRIES"
            ws[f"A{row}"].fill = _fill(COLORS["SubHeader"])
            ws[f"A{row}"].font = Font(bold=True, color="FFFFFF", size=11)
            ws[f"A{row}"].alignment = _align("center")
            row += 1
            for std in vcm.key_standards:
                ws.merge_cells(f"A{row}:D{row}")
                ws[f"A{row}"].value = f"• {std}"
                ws[f"A{row}"].alignment = _align("left")
                ws[f"A{row}"].fill = _fill(COLORS["LightGrey"])
                ws[f"A{row}"].border = _border()
                row += 1

        # Insight
        if vcm.insight and vcm.insight != "missing":
            row += 1
            ws.merge_cells(f"A{row}:D{row}")
            ws[f"A{row}"].value = f"💡 Insight: {vcm.insight}"
            ws[f"A{row}"].alignment = _align("left", wrap=True)
            ws[f"A{row}"].fill = _fill(COLORS["LightBlue"])
            ws[f"A{row}"].border = _border()
            ws.row_dimensions[row].height = 60

        col_widths = [25, 20, 18, 35]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------
    # Sheet: Carbon Offsets
    # ------------------------------------------------------------------

    def _add_offsets_sheet(self, wb: Workbook, report: CarbonIntelligenceReport) -> None:
        ws = wb.create_sheet("Carbon Offsets")
        ws.sheet_view.showGridLines = False

        self._sheet_title(ws, "CARBON OFFSET PROJECT QUALITY ASSESSMENT", 8)

        headers = [
            "Project Type", "Region", "Standard", "Price Range",
            "Additionality", "Permanence", "Quality", "Signal"
        ]
        row = self._write_headers(ws, headers, start_row=3)

        if not report.carbon_offsets:
            ws.cell(row=row + 1, column=1, value="No carbon offset data available.")
        else:
            for i, offset in enumerate(report.carbon_offsets):
                row += 1
                bg = COLORS["LightBlue"] if i % 2 == 0 else COLORS["White"]
                vals = [
                    offset.project_type, offset.region, offset.standard,
                    f"{offset.price_range} {offset.currency}".strip(),
                    offset.additionality, offset.permanence,
                    offset.quality_assessment, offset.color_code,
                ]
                for col, val in enumerate(vals, 1):
                    cell = ws.cell(row=row, column=col, value=val)
                    cell.fill = _fill(bg)
                    cell.alignment = _align()
                    cell.border = _border()
                    if col == 7 and val in ("High", "Medium", "Low"):
                        cell.fill = _fill(_perf_color(val))
                        cell.font = Font(bold=True, color="FFFFFF", size=10)
                    if col == 8 and val in ("Green", "Yellow", "Red"):
                        cell.fill = _fill(_color_code_hex(val))
                        cell.value = f"● {val}"
                        cell.font = Font(bold=True, color="FFFFFF", size=10)

        col_widths = [22, 14, 16, 18, 14, 12, 10, 10]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------
    # Sheet: Risk Analysis
    # ------------------------------------------------------------------

    def _add_risk_sheet(self, wb: Workbook, report: CarbonIntelligenceReport) -> None:
        ws = wb.create_sheet("Risk Analysis")
        ws.sheet_view.showGridLines = False

        self._sheet_title(ws, "CARBON MARKET RISK ANALYSIS", 3)

        risk = report.risk_analysis
        row = 3

        # Overall risk badge
        row += 1
        ws.merge_cells(f"A{row}:C{row}")
        ws[f"A{row}"].value = f"OVERALL RISK LEVEL: {risk.overall_risk_level.upper()}"
        hex_color = _color_code_hex(risk.color_code)
        ws[f"A{row}"].fill = _fill(hex_color)
        ws[f"A{row}"].font = Font(bold=True, color="FFFFFF", size=14)
        ws[f"A{row}"].alignment = _align("center", wrap=False)
        ws.row_dimensions[row].height = 30
        row += 2

        risk_categories = [
            ("MARKET RISKS", risk.market_risks, COLORS["Red"]),
            ("PRICING RISKS", risk.pricing_risks, COLORS["Yellow"]),
            ("POLICY RISKS", risk.policy_risks, COLORS["SubHeader"]),
        ]

        for category_name, risk_list, color in risk_categories:
            ws.merge_cells(f"A{row}:C{row}")
            ws[f"A{row}"].value = category_name
            ws[f"A{row}"].fill = _fill(color)
            ws[f"A{row}"].font = Font(bold=True, color="FFFFFF", size=11)
            ws[f"A{row}"].alignment = _align("center")
            row += 1
            if risk_list:
                for item in risk_list:
                    ws.merge_cells(f"A{row}:C{row}")
                    ws[f"A{row}"].value = f"⚠ {item}"
                    ws[f"A{row}"].alignment = _align("left", wrap=True)
                    ws[f"A{row}"].fill = _fill(COLORS["LightGrey"])
                    ws[f"A{row}"].border = _border()
                    ws.row_dimensions[row].height = 30
                    row += 1
            else:
                ws.merge_cells(f"A{row}:C{row}")
                ws[f"A{row}"].value = "No risks identified."
                ws[f"A{row}"].fill = _fill(COLORS["White"])
                ws[f"A{row}"].border = _border()
                row += 1
            row += 1

        if risk.insight and risk.insight != "missing":
            ws.merge_cells(f"A{row}:C{row}")
            ws[f"A{row}"].value = f"💡 {risk.insight}"
            ws[f"A{row}"].alignment = _align("left", wrap=True)
            ws[f"A{row}"].fill = _fill(COLORS["LightBlue"])
            ws[f"A{row}"].border = _border()
            ws.row_dimensions[row].height = 60

        col_widths = [30, 30, 30]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------
    # Sheet: Opportunities & Forecast
    # ------------------------------------------------------------------

    def _add_opportunities_sheet(self, wb: Workbook, report: CarbonIntelligenceReport) -> None:
        ws = wb.create_sheet("Opportunities & Forecast")
        ws.sheet_view.showGridLines = False

        self._sheet_title(ws, "INVESTMENT OPPORTUNITIES & PRICE FORECAST", 3)

        opp = report.opportunity_analysis
        forecast = report.forecast
        row = 3

        # --- Opportunities ---
        opp_sections = [
            ("💰 INVESTMENT OPPORTUNITIES", opp.investment_opportunities, COLORS["Green"]),
            ("🏢 CORPORATE STRATEGY", opp.corporate_strategy_opportunities, COLORS["SubHeader"]),
            ("🌱 ESG OPPORTUNITIES", opp.esg_opportunities, COLORS["Gold"]),
        ]
        for section_name, items, color in opp_sections:
            row += 1
            ws.merge_cells(f"A{row}:C{row}")
            ws[f"A{row}"].value = section_name
            ws[f"A{row}"].fill = _fill(color)
            ws[f"A{row}"].font = Font(bold=True, color="FFFFFF", size=11)
            ws[f"A{row}"].alignment = _align("center")
            row += 1
            if items:
                for item in items:
                    ws.merge_cells(f"A{row}:C{row}")
                    ws[f"A{row}"].value = f"✓ {item}"
                    ws[f"A{row}"].alignment = _align("left", wrap=True)
                    ws[f"A{row}"].fill = _fill(COLORS["LightGrey"])
                    ws[f"A{row}"].border = _border()
                    ws.row_dimensions[row].height = 28
                    row += 1
            else:
                ws.merge_cells(f"A{row}:C{row}")
                ws[f"A{row}"].value = "No specific opportunities identified."
                ws[f"A{row}"].border = _border()
                row += 1
            row += 1

        # --- Forecast ---
        row += 1
        ws.merge_cells(f"A{row}:C{row}")
        ws[f"A{row}"].value = "📈 PRICE FORECAST"
        ws[f"A{row}"].fill = _fill(COLORS["Header"])
        ws[f"A{row}"].font = Font(bold=True, color="FFFFFF", size=12)
        ws[f"A{row}"].alignment = _align("center")

        forecast_rows = [
            ("Short-Term Outlook (1-2 yr)", forecast.short_term_outlook),
            ("Short-Term Price Target", forecast.price_target_short_term),
            ("Long-Term Outlook (5-10 yr)", forecast.long_term_outlook),
            ("Long-Term Price Target", forecast.price_target_long_term),
            ("Price Direction", forecast.price_direction),
            ("Forecast Confidence", forecast.confidence),
        ]
        for label, value in forecast_rows:
            row += 1
            ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=10)
            ws.cell(row=row, column=1).border = _border()
            ws.cell(row=row, column=1).fill = _fill(COLORS["LightBlue"])
            value_cell = ws.cell(row=row, column=2, value=value)
            value_cell.alignment = _align("left", wrap=True)
            value_cell.border = _border()
            ws.merge_cells(f"B{row}:C{row}")
            ws.row_dimensions[row].height = 22

        col_widths = [35, 30, 25]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------
    # Utility methods
    # ------------------------------------------------------------------

    def _sheet_title(self, ws: Any, title: str, col_span: int) -> None:
        """Write a centred title row at row 1."""
        col_letter = get_column_letter(col_span)
        ws.merge_cells(f"A1:{col_letter}1")
        cell = ws["A1"]
        cell.value = title
        cell.fill = _fill(COLORS["Header"])
        cell.font = Font(bold=True, color="FFFFFF", size=14)
        cell.alignment = _align("center", wrap=False)
        ws.row_dimensions[1].height = 32

    def _write_headers(
        self,
        ws: Any,
        headers: list[str],
        start_row: int,
        col_count: int | None = None,
    ) -> int:
        """Write a header row and return the row number."""
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col, value=h)
            cell.fill = _fill(COLORS["SubHeader"])
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.alignment = _align("center", wrap=False)
            cell.border = _border()
        ws.row_dimensions[start_row].height = 22
        return start_row
